"""Restitue les originaux, construit les documents et vérifie les fichiers livrés."""
from pathlib import Path
import base64,hashlib,io,json,os,subprocess,sys,tarfile,tempfile,shutil
ROOT=Path(os.environ.get('FORMATION_ROOT',str(Path.cwd()/'formation_4_020'))).resolve()
TMP=Path(os.environ.get('PRODUCTION_TEMP',tempfile.mkdtemp(prefix='p2e4020_'))).resolve();TMP.mkdir(parents=True,exist_ok=True)
os.environ['FORMATION_ROOT']=str(ROOT);os.environ['PRODUCTION_TEMP']=str(TMP)
REPO=ROOT.parent

def extract_originals():
 parts=sorted((REPO/'.publication_4020/originaux').glob('part_*'))
 if parts:
  raw=b''.join(p.read_bytes() for p in parts)
  assert len(raw)==136640 and hashlib.sha256(raw).hexdigest()=='69032d1f9289798fbe7afdc35965ab754ed6396df9b0678b8d12114c455eb94b','Archive originale invalide'
  with tarfile.open(fileobj=io.BytesIO(raw),mode='r:xz') as archive:
   members=archive.getmembers();assert len(members)==13
   for m in members:
    assert m.isfile() and m.name.startswith('formation_4_020/') and '..' not in Path(m.name).parts
    data=archive.extractfile(m).read();dest=(REPO/m.name).resolve();assert dest.is_relative_to(ROOT)
    if dest.exists() and dest.read_bytes()!=data:raise RuntimeError('Original existant différent : '+str(dest))
    dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(data)
 return {str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.rglob('*') if p.is_file() and (p.parent.name=='travaux_pratiques' or p.name=='contenus_diaporamas.json')}

def cmd(args,timeout=120):
 p=subprocess.run(args,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=timeout)
 print(p.stdout,flush=True)
 if p.returncode:raise RuntimeError('Commande échouée : '+repr(args))

def convert(src,out,profile):
 cmd(['libreoffice','-env:UserInstallation=file://'+str(TMP/profile),'--headless','--convert-to','pdf','--outdir',str(out),str(src)],90)
 dst=out/(src.stem+'.pdf')
 if not dst.exists() or dst.stat().st_size<1000:raise RuntimeError('Export PDF absent : '+str(dst))
 return dst

def export_slides():
 from pptx import Presentation
 import fitz
 for day,count in [(1,75),(2,83)]:
  src=ROOT/f'presentations/Formation_4_020_Jour_{day}.pptx';dest=src.with_suffix('.pdf');parts=TMP/f'jour{day}';parts.mkdir(exist_ok=True);pdfs=[]
  for start in range(0,count,10):
   end=min(count,start+10);r=Presentation(src)
   for i in reversed(range(len(r.slides))):
    if not(start<=i<end):sid=r.slides._sldIdLst[i];r.part.drop_rel(sid.rId);r.slides._sldIdLst.remove(sid)
   part=parts/f'part_{start:03}.pptx';r.save(part);pdfs.append(convert(part,parts,f'lo_j{day}_{start}'))
  merged=fitz.open()
  for p in pdfs:
   with fitz.open(p) as d:merged.insert_pdf(d)
  assert len(merged)==count
  merged.save(dest,garbage=4,deflate=True)

def validate(originals):
 from pptx import Presentation
 from docx import Document
 import fitz
 report={'originaux':{},'presentations':{},'documents':{},'illustrations':len(list((ROOT/'illustrations').glob('*.png')))}
 for name,sha in originals.items():
  p=ROOT/name;assert hashlib.sha256(p.read_bytes()).hexdigest()==sha;report['originaux'][name]=sha
 assert len(originals)==13
 for day,count in [(1,75),(2,83)]:
  p=ROOT/f'presentations/Formation_4_020_Jour_{day}.pptx';r=Presentation(p);assert len(r.slides)==count
  assert all(len(s.notes_slide.notes_text_frame.text)>100 for s in r.slides)
  with fitz.open(p.with_suffix('.pdf')) as d:
   assert len(d)==count;assert all(len(q.get_text().strip())>40 for q in d)
  report['presentations'][p.name]={'diapositives':count,'notes':count,'pdf_pages':count}
 for name in ['Dossier_pedagogique_4_020','Cahier_participant_4_020']:
  p=ROOT/f'documents/{name}.docx';doc=Document(p);assert len(doc.paragraphs)>100
  with fitz.open(p.with_suffix('.pdf')) as d:pages=len(d);words=sum(len(q.get_text().split()) for q in d)
  assert pages>=10;report['documents'][p.name]={'pages_pdf':pages,'mots_pdf':words}
 report['fichiers']={str(p.relative_to(ROOT)):{'octets':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in ROOT.rglob('*') if p.is_file() and p.name!='MANIFESTE_LIVRAISON.json'}
 (ROOT/'MANIFESTE_LIVRAISON.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
 return report

def readme(report):
 g=report['documents']['Dossier_pedagogique_4_020.docx']['pages_pdf'];c=report['documents']['Cahier_participant_4_020.docx']['pages_pdf']
 t=f'''# Formation 4-020 : intelligence artificielle et systèmes d’information

Lot 4. Niveau Application (A). 12 heures sur 2 jours. Aucun client ni date de session attribué.

## Supports à ouvrir

| Support | Version modifiable | PDF |
| --- | --- | --- |
| Jour 1, 75 diapositives avec notes | [PowerPoint](presentations/Formation_4_020_Jour_1.pptx) | [PDF](presentations/Formation_4_020_Jour_1.pdf) |
| Jour 2, 83 diapositives avec notes | [PowerPoint](presentations/Formation_4_020_Jour_2.pptx) | [PDF](presentations/Formation_4_020_Jour_2.pdf) |
| Dossier pédagogique et guide formateur, {g} pages | [Word](documents/Dossier_pedagogique_4_020.docx) | [PDF](documents/Dossier_pedagogique_4_020.pdf) |
| Cahier participant, {c} pages | [Word](documents/Cahier_participant_4_020.docx) | [PDF](documents/Cahier_participant_4_020.pdf) |

## Travaux pratiques

Le dossier `travaux_pratiques` conserve les 12 fichiers originaux : trois notebooks, données synthétiques Novalia, résultats de référence, instructions et dépendances. Commencer par [LISEZ_MOI.md](travaux_pratiques/LISEZ_MOI.md).

Le conducteur totalise 720 minutes : 450 minutes de TP effectifs (62,5 %), 210 minutes d’apports et 60 minutes de diagnostic et d’évaluation. Les annexes ne constituent pas des heures contractuelles supplémentaires.

## Contenu et vérification

Les PowerPoint contiennent des schémas vectoriels éditables et des notes formateur. Les figures du dossier sont issues de ces schémas. Aucun visuel ImageGen n’est utilisé. Aucun fichier PBIX n’est inclus ; les données, consignes et maquettes de restitution Power BI sont fournies.

Les cas commerciaux distinguent faits documentés, déclarations fournisseur et simulations. Le corpus est inclus dans le guide. Les pages réglementaires doivent être revérifiées avant une future session.

Le fichier [MANIFESTE_LIVRAISON.json](MANIFESTE_LIVRAISON.json) contient les comptes de pages, les contrôles structurels et les empreintes SHA-256 des fichiers. Les 13 fichiers initiaux sont vérifiés séparément et conservés sans altération.

Les scripts de fabrication sont dans `sources_production`. Ils ne sont pas nécessaires pour lire les supports livrés.

https://p2enjoy.studio
'''
 (ROOT/'README.md').write_text(t,encoding='utf-8')

if __name__=='__main__':
 originals=extract_originals()
 src=ROOT/'sources_production';cmd([sys.executable,str(src/'build_slides.py')]);export_slides();cmd([sys.executable,str(src/'build_documents.py')])
 for name in ['Dossier_pedagogique_4_020','Cahier_participant_4_020']:convert(ROOT/f'documents/{name}.docx',ROOT/'documents','lo_'+name)
 if (src/'export_markdown.py').exists():cmd([sys.executable,str(src/'export_markdown.py')])
 report=validate(originals);readme(report);validate(originals)
 print('LIVRAISON_LOCALE_VERIFIEE : 13 originaux, 2 PowerPoint, 2 Word, 4 PDF',flush=True)
