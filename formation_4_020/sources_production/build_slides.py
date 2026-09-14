from pathlib import Path
import json,math,re,os,textwrap
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from PIL import ImageFont
ROOT=Path(os.environ.get('FORMATION_ROOT','/mnt/data/formation_4_020'))
D=json.loads((ROOT/'sources_production/contenus_diaporamas.json').read_text())
OUT=ROOT/'presentations';OUT.mkdir(exist_ok=True)
C={'blue':'23468C','green':'238C33','yellow':'D9CF4A','red':'F24141','ink':'0D0D0D','muted':'586777','line':'D9E1EB','pale':'F1F5FA','white':'FFFFFF','navy':'132443'}
FONT='Liberation Sans';FP='/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'
if not Path(FP).exists():FP='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
W,H=13.333333,7.5
ACC=[C['blue'],C['green'],C['blue'],C['green'],C['blue'],C['green']]
META=[]

def clean(v):
 if v is None:return ''
 return str(v).replace('—',': ').replace('–',', ').replace('\\(','').replace('\\)','').replace('\\[','').replace('\\]','')
def fit(text,w,h,size=23,minimum=13,bold=False):
 text=clean(text); size=float(size)
 while size>=minimum:
  font_path=FP.replace('-Regular.ttf','-Bold.ttf') if bold else FP
  if not Path(font_path).exists():font_path=FP
  f=ImageFont.truetype(font_path,max(1,int(size*1.5)))
  if any(f.getlength(word)>w*108*.88 for word in text.split()):
   size-=.5;continue
  count=0
  for para in text.split('\n'):
   words=para.split();line=''
   if not words:count+=1;continue
   for word in words:
    candidate=line+' '+word if line else word
    if f.getlength(candidate)>w*108*.91 and line:count+=1;line=word
    else:line=candidate
   count+=1
  if count*size*1.20<=h*72*.94:return size
  size-=.5
 return minimum

def shape(s,typ,x,y,w,h,fill=None,line=None,trans=0):
 q=s.shapes.add_shape(typ,Inches(x),Inches(y),Inches(w),Inches(h))
 if fill:q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 else:q.fill.background()
 q.line.fill.background() if not line else None
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(.8)
 return q

def txt(s,t,x,y,w,h,sz=23,col=None,bold=False,align=None,val=MSO_ANCHOR.MIDDLE,minimum=13):
 t=clean(t);q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=Inches(.025);tf.margin_right=Inches(.025);tf.margin_top=Inches(.025);tf.margin_bottom=Inches(.025);tf.vertical_anchor=val
 size=fit(t,w,h,sz,minimum,bold)
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(col or C['ink']);p.space_after=Pt(3);p.space_before=Pt(0)
  if align is not None:p.alignment=align
 META.append({'slide':CUR,'text':t,'x':x,'y':y,'w':w,'h':h,'size':size})
 return q

def line(s,x1,y1,x2,y2,col=None,width=2):
 c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));c.line.color.rgb=RGBColor.from_string(col or C['line']);c.line.width=Pt(width);return c

def card(s,title,body,x,y,w,h,accent=None,num=None,dark=False):
 ac=accent or C['blue'];shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,C['navy'] if dark else C['pale'])
 shape(s,MSO_SHAPE.RECTANGLE,x,y,.06,h,ac)
 if h<1.30:
  txt(s,title,x+.17,y+.10,w-.34,.30,18,C['white'] if dark else ac,True,minimum=12)
  txt(s,body,x+.17,y+.44,w-.34,max(.20,h-.50),16,C['white'] if dark else C['ink'],val=MSO_ANCHOR.TOP,minimum=11)
  return
 if num:
  txt(s,str(num),x+.19,y+.14,.55,.31,15,ac,True)
  th=.77;ty=y+.57;by=y+1.43;bh=h-1.60;ts=21
 else:
  th=.72 if len(clean(title))>30 else .52;ty=y+.15;by=y+.22+th;bh=h-th-.38;ts=23
 txt(s,title,x+.19,ty,w-.38,th,ts,C['white'] if dark else ac,True,minimum=13)
 txt(s,body,x+.18,by,w-.36,max(.3,bh),21,C['white'] if dark else C['ink'],val=MSO_ANCHOR.TOP)

def pairs(items):
 return [(clean(a[0]),clean('\n'.join(map(str,a[1:])))) if isinstance(a,list) else (clean(a),'') for a in items]

def grid(s,it,y=1.95,h=4.25,cols=None):
 a=pairs(it);n=len(a);cols=cols or (2 if n<=4 else 3);rows=math.ceil(n/cols);gap=.22;cw=(12.03-gap*(cols-1))/cols;ch=(h-gap*(rows-1))/rows
 for i,(t,b) in enumerate(a):card(s,t,b,.65+(i%cols)*(cw+gap),y+(i//cols)*(ch+gap),cw,ch,ACC[i%len(ACC)])

def flow(s,it,y=2.18,h=3.25):
 a=pairs(it);n=len(a)
 if n>6:return grid(s,it)
 gap=.31;cw=(12.03-gap*(n-1))/n
 for i,(t,b) in enumerate(a):
  x=.65+i*(cw+gap);card(s,t,b,x,y,cw,h,ACC[i%6],f'{i+1:02}')
  if i<n-1:shape(s,MSO_SHAPE.CHEVRON,x+cw+.055,y+h/2-.13,.20,.26,C['green'])

def table(s,it):
 rows=it;nr=len(rows);nc=max(map(len,rows));y=1.96;rh=min(.75,4.30/nr);width=12.03
 colweights=([.29,.35,.36] if nc==3 else [.31]+[(.69)/(nc-1)]*(nc-1))
 if nc==2:colweights=[.35,.65]
 for i,row in enumerate(rows):
  x=.65
  for j,val in enumerate(row):
   w=width*colweights[j];shape(s,MSO_SHAPE.RECTANGLE,x,y+i*rh,w,rh,C['blue'] if i==0 else (C['pale'] if i%2 else C['white']))
   txt(s,val,x+.12,y+i*rh+.045,w-.24,rh-.09,21 if nc<=3 else 18,C['white'] if i==0 else C['ink'],i==0,minimum=13)
   x+=w
  line(s,.65,y+(i+1)*rh,12.68,y+(i+1)*rh,C['line'],.6)

def frame(s,d):
 s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(C['white'])
 shape(s,MSO_SHAPE.RECTANGLE,0,0,.13,H,C['blue'] if d['day']==1 else C['green'])
 txt(s,f"FORMATION 4-020   /   JOUR {d['day']}   /   {d['block']}",.65,.28,8,.28,11,C['muted'],True)
 mode=d.get('mode','');shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,10.70,.26,1.98,.35,C['pale']);txt(s,mode.upper(),10.78,.28,1.82,.29,10,C['blue'],True,PP_ALIGN.CENTER)
 title=clean(d['title']).replace('\n',' ');txt(s,title,.65,.86,12.03,.88,33,C['navy'],True,minimum=24)
 shape(s,MSO_SHAPE.RECTANGLE,.65,6.45,12.03,.52,C['navy']);txt(s,d.get('take',''),.82,6.52,11.65,.36,16,C['white'],True,minimum=12)
 txt(s,'P2ENJOY  ·  https://p2enjoy.studio',.65,7.15,7,.18,9,C['muted'])
 refs=' · '.join(d.get('sources',[]));txt(s,refs,7.10,7.15,4.50,.18,8,C['muted'],align=PP_ALIGN.RIGHT)
 txt(s,d['id'],11.85,7.10,.84,.28,10,C['blue'],True,PP_ALIGN.RIGHT)

def cover(s,d):
 s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(C['navy'])
 shape(s,MSO_SHAPE.RECTANGLE,0,0,.18,H,C['green'])
 txt(s,'P2ENJOY  /  FORMATION 4-020',.65,.42,10,.35,13,C['white'],True)
 txt(s,'JOUR '+str(d['day']),.65,1.40,8,.54,20,C['yellow'],True)
 txt(s,d['title'],.65,2.02,9.70,2.20,43,C['white'],True,minimum=33)
 txt(s,d.get('subtitle',''),.65,4.42,10, .65,25,C['white'])
 labs=pairs(d['items']);t='   ·   '.join(t+((' : '+b) if b else '') for t,b in labs)
 txt(s,t,.65,5.48,10.8,.7,20,C['yellow'])
 # Native information network at right
 for i,(x,y) in enumerate([(11.2,1.35),(11.65,2.55),(10.85,3.70),(11.7,4.85)]):
  if i:line(s,px+.28,py+.28,x+.28,y+.28,C['green'],3)
  shape(s,MSO_SHAPE.HEXAGON,x,y,.56,.56,C['green'] if i%2 else C['blue']);px,py=x,y
 txt(s,d['take'],.65,6.60,11.8,.35,14,C['white']);txt(s,'https://p2enjoy.studio',.65,7.11,10,.22,10,C['white'])

def section(s,d):
 s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(C['blue'] if d['day']==1 else C['navy'])
 txt(s,f"JOUR {d['day']}  /  {d['block']}  /  NIVEAU APPLICATION",.65,.50,11,.4,14,C['white'],True)
 txt(s,d['title'],.65,1.48,11.6,2.15,42,C['white'],True,minimum=30)
 for i,(t,b) in enumerate(pairs(d['items'])):txt(s,f'{i+1:02}  {t}'+(' : '+b if b else ''),.72,4.02+i*.60,11.8,.5,24,C['white'])
 txt(s,d['take'],.65,6.65,12,.42,18,C['yellow'],True);txt(s,'https://p2enjoy.studio',.65,7.17,10,.16,9,C['white'])

def hubs(s,d):
 a=pairs(d['items']);n=len(a);cx,cy=6.65,4.15
 positions=[(cx+4.12*math.cos(-math.pi/2+2*math.pi*i/n),cy+1.57*math.sin(-math.pi/2+2*math.pi*i/n)) for i in range(n)]
 if d['kind']=='cycle':
  for i,((x,y),(xx,yy)) in enumerate(zip(positions,positions[1:]+positions[:1])):
   line(s,x,y,xx,yy,C['green'],2)
   arrow=shape(s,MSO_SHAPE.CHEVRON,(x+xx)/2-.11,(y+yy)/2-.08,.22,.16,C['green']);arrow.rotation=math.degrees(math.atan2(yy-y,xx-x))
 for i,(t,b) in enumerate(a):
  x,y=positions[i]
  if d['kind']!='cycle':line(s,cx,cy,x,y,C['line'],2)
  shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x-1.27,y-.44,2.54,.88,C['pale']);txt(s,t,x-1.18,y-.38,2.36,.36,19,C['blue'],True,PP_ALIGN.CENTER);txt(s,b,x-1.18,y-.04,2.36,.36,16,align=PP_ALIGN.CENTER)
 shape(s,MSO_SHAPE.OVAL,cx-1.1,cy-.72,2.2,1.44,C['blue']);txt(s,d.get('center','DÉCISION\nMÉTIER'),cx-.99,cy-.61,1.98,1.22,23,C['white'],True,PP_ALIGN.CENTER)

def pyramid(s,d):
 a=pairs(d['items']);n=len(a)
 for i,(t,b) in enumerate(a):
  w=4.2+i*1.34;x=6.66-w/2;y=2.02+i*.93
  shape(s,MSO_SHAPE.TRAPEZOID,x,y,w,.83,ACC[(n-1-i)%6]);txt(s,t,x+.50,y+.10,w-1,.30,21,C['white'],True,PP_ALIGN.CENTER);txt(s,b,x+.50,y+.42,w-1,.28,16,C['white'],align=PP_ALIGN.CENTER)

def networks(s,d):
 a=pairs(d['items']);kind=d['kind']
 if kind=='architecture':
  flow(s,d['items'],2.45,2.9);line(s,1.55,5.75,11.8,5.75,C['green'],2);txt(s,'Identités  •  Contrôles  •  Observabilité  •  Responsabilités',2,5.8,9.5,.33,17,C['green'],True,PP_ALIGN.CENTER);return
 if kind in ['novalia','network','open_si','si','production','security']: 
  if len(a)<=4:flow(s,d['items']);return
  grid(s,d['items'],y=2.12,h=3.7,cols=3);return
 hubs(s,d)

def plot_native(s,d):
 # Each conceptual chart is made of native vector elements, not screenshots.
 kind=d['kind'];x0,y0,w,h=.90,2.08,7.5,3.63
 line(s,x0,y0+h,x0+w,y0+h,C['muted'],1.1);line(s,x0,y0,x0,y0+h,C['muted'],1.1)
 def xy(x,y):return x0+x*w,y0+h-y*h
 if kind=='overfit':
  # Three mini-panels; generic main axes removed for clean composition
  for sh in list(s.shapes)[-2:]:sh._element.getparent().remove(sh._element)
  for j,(t,b) in enumerate(pairs(d['items'])):
   xx=.82+j*4.13;yy=2.25;ww=3.65;hh=2.50
   line(s,xx,yy+hh,xx+ww,yy+hh,C['line'],1);line(s,xx,yy,xx,yy+hh,C['line'],1)
   points=[]
   for i in range(11):
    u=(i+.5)/12;v=.15+.65*u+.10*math.sin(i*3.8)
    shape(s,MSO_SHAPE.OVAL,xx+u*ww-.045,yy+(1-v)*hh-.045,.09,.09,C['muted']);points.append((u,v))
   old=None
   for i in range(75):
    u=i/74
    v=.40 if j==0 else .15+.65*u if j==1 else .15+.65*u+.15*math.sin(u*37)
    px=xx+u*ww;py=yy+(1-v)*hh
    if old:line(s,*old,px,py,C['red'] if j!=1 else C['green'],2)
    old=(px,py)
   txt(s,t,xx,5.03,ww,.45,24,C['blue'],True,PP_ALIGN.CENTER);txt(s,b,xx,5.52,ww,.52,19,align=PP_ALIGN.CENTER)
  return
 if kind=='logistic':vals=[(i/90,1/(1+math.exp(-10*(i/90-.5)))) for i in range(91)]
 elif kind=='drift':vals=[(i/90,.85-.55*(i/90)**1.9+.025*math.sin(i)) for i in range(91)]
 elif kind=='threshold':vals=[(i/90,1-i/90) for i in range(91)]
 else:vals=[(i/90,.30+.22*math.sin(i/15)) for i in range(91)]
 for (u,v),(u2,v2) in zip(vals,vals[1:]):line(s,*xy(u,v),*xy(u2,v2),C['blue'],2.6)
 if kind=='threshold':
  for i in range(90):line(s,*xy(i/90,i/90),*xy((i+1)/90,(i+1)/90),C['green'],2.6)
 txt(s,'Illustration conceptuelle, pas une mesure du TP',1,5.90,7.3,.27,11,C['muted'])
 a=pairs(d['items']);ch=3.8/max(1,len(a))
 for i,(t,b) in enumerate(a):card(s,t,b,8.83,2.03+i*ch,3.80,ch-.10,ACC[i%6])

def neural(s,d):
 layers=[4,5,4,2];xs=[1.5,3.60,5.70,7.80];nodes=[]
 for k,(x,n) in enumerate(zip(xs,layers)):
  ys=[2.55+i*(2.40/(n-1)) for i in range(n)];nodes.append([(x,y) for y in ys])
 for a,b in zip(nodes,nodes[1:]):
  for x,y in a:
   for xx,yy in b:line(s,x,y,xx,yy,'CED9E9',.7)
 for k,a in enumerate(nodes):
  for x,y in a:shape(s,MSO_SHAPE.OVAL,x-.15,y-.15,.3,.3,C['green'] if k==len(nodes)-1 else C['blue'])
 txt(s,'Entrées',.8,5.25,1.5,.4,18,C['blue'],True,PP_ALIGN.CENTER);txt(s,'Représentations apprises',2.7,5.25,4.2,.4,18,C['blue'],True,PP_ALIGN.CENTER);txt(s,'Sortie',7.1,5.25,1.6,.4,18,C['green'],True,PP_ALIGN.CENTER)
 a=pairs(d['items'])
 for i,(t,b) in enumerate(a):card(s,t,b,9,2.05+i*(3.95/len(a)),3.68,3.95/len(a)-.13,ACC[i%6])

def tree(s,d):
 coords=[(5.1,2.08,3.0,.67,'Satisfaction faible ?'),(2.1,3.52,3.0,.65,'Tickets nombreux ?'),(8.5,3.52,3.0,.65,'Risque modéré'),(.8,5.0,2.8,.68,'Risque élevé'),(4.3,5.0,2.8,.68,'Risque à examiner')]
 for a,b in [(0,1),(0,2),(1,3),(1,4)]:
  x,y,w,h,_=coords[a];xx,yy,ww,hh,_=coords[b];line(s,x+w/2,y+h,xx+ww/2,yy,C['line'],2)
 for i,(x,y,w,h,t) in enumerate(coords):shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,C['blue'] if i<2 else C['pale']);txt(s,t,x+.10,y+.05,w-.2,h-.1,20,C['white'] if i<2 else C['green'],True,PP_ALIGN.CENTER)
 txt(s,'Exemple schématique. Les seuils réels sont appris et validés.',.90,5.91,11.8,.30,14,C['muted'],align=PP_ALIGN.CENTER)

def dashboard(s,d):
 a=pairs(d['items'])
 for i,(t,b) in enumerate(a):card(s,t,b,.65+i*3.06,2.00,2.83,1.14,ACC[i%6])
 for xx,ww,label in [(.65,5.9,'Répartition des scores'),(6.77,5.91,'Alertes par segment')]:
  shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,xx,3.38,ww,2.58,C['pale']);txt(s,label,xx+.18,3.53,ww-.36,.35,21,C['blue'],True)
  for j,v in enumerate([.3,.55,.90,.72,.35,.20,.1]):shape(s,MSO_SHAPE.RECTANGLE,xx+.45+j*.64,5.52-v*1.30,.43,v*1.30,C['blue'] if xx<3 else C['green'])
 txt(s,'Maquette illustrative. Les valeurs du tableau réel proviennent du fichier CSV.',1,6.06,11.6,.24,12,C['muted'],align=PP_ALIGN.CENTER)

def special(s,d):
 k=d['kind'];it=d['items'];a=pairs(it)
 if k in ('cards','case','comparison','gates','canvas','activity','risk','house','iceberg','funnel','clusters','attention','notebook'):
  if k=='activity':
   grid(s,it,y=2.15,h=3.7);txt(s,f"ATELIER  /  {d['segment_minutes']} MINUTES POUR LA SÉQUENCE",.69,1.81,11,.25,11,C['green'],True)
  elif k=='house':
   shape(s,MSO_SHAPE.CHEVRON,.8,2.1,11.7,.66,C['blue']);txt(s,'CONFIANCE DANS LES DÉCISIONS',1.1,2.20,11.1,.38,23,C['white'],True,PP_ALIGN.CENTER);grid(s,it,y=3.02,h=2.92,cols=5)
  elif k=='funnel':
   for i,(t,b) in enumerate(a):
    ww=11.8-i*1.65;xx=(13.333-ww)/2;yy=2+i*.96;shape(s,MSO_SHAPE.TRAPEZOID,xx,yy,ww,.84,ACC[i%6]);txt(s,t+' : '+b,xx+.55,yy+.16,ww-1.1,.52,23,C['white'],True,PP_ALIGN.CENTER)
  else:grid(s,it,cols=2 if k in ['case','comparison','risk'] else None)
 elif k in ('flow','steps','train','timeline','agenda','wide','api','latency','split','neuron'):
  if k=='agenda':
   for i,(t,b) in enumerate(a):txt(s,t,.85,2.02+i*.78,2.15,.62,21,C['blue'],True);txt(s,b,3.1,2.02+i*.78,9.3,.62,24);line(s,.85,2.73+i*.78,12.45,2.73+i*.78,C['line'],.7)
  elif k=='split':
   widths=[7.15,2.36,2.36];xx=.65
   for i,(t,b) in enumerate(a):shape(s,MSO_SHAPE.RECTANGLE,xx,2.3,widths[i],1.2,ACC[i%6]);txt(s,t,xx+.15,2.5,widths[i]-.3,.65,35,C['white'],True,PP_ALIGN.CENTER);txt(s,b,xx+.10,3.80,widths[i]-.2,1.15,26,C['navy'],True,PP_ALIGN.CENTER);xx+=widths[i]+.09
  else:flow(s,it)
 elif k in ('table','confusion'):table(s,it)
 elif k=='pyramid':pyramid(s,d)
 elif k in ('hub','cycle'):hubs(s,d)
 elif k=='nested':
  for i,(t,b) in enumerate(a[:3]):x=.75+i*.55;y=2.0+i*1.02;w=8.1-i*.82;h=4.05-i*.95;shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,[C['pale'],'DCE6F8',C['blue']][i]);txt(s,t,x+.23,y+.12,w-.46,.38,25,C['white'] if i==2 else C['blue'],True);txt(s,b,x+.23,y+.54,w-.46,.35,18,C['white'] if i==2 else C['ink'])
  card(s,*a[3],9.2,2.57,3.35,2.72,C['green'])
 elif k in ('si','novalia','production','architecture','open_si','network','security'):networks(s,d)
 elif k in ('overfit','logistic','threshold','drift'):plot_native(s,d)
 elif k=='neural':neural(s,d)
 elif k=='tree':tree(s,d)
 elif k=='dashboard':dashboard(s,d)
 elif k=='metric':
  txt(s,d.get('metric',''),.8,2.15,6,1.65,95,C['blue'],True,minimum=65);txt(s,d.get('label',''),.9,3.92,5.6,1.15,28,C['navy'],True)
  for i,(t,b) in enumerate(a):card(s,t,b,7.20,2.04+i*(3.90/len(a)),5.48,3.90/len(a)-.13,ACC[i%6])
 elif k=='quiz':
  for i,(t,b) in enumerate(a):
   yy=2.00+i*(4.02/len(a));shape(s,MSO_SHAPE.OVAL,.78,yy+.03,.47,.47,C['blue']);txt(s,t if len(t)<=2 else str(i+1),.79,yy+.07,.45,.30,15,C['white'],True,PP_ALIGN.CENTER);txt(s,b,1.47,yy,11.08,4.02/len(a)-.12,25,minimum=18)
 elif k=='code':
  code=d.get('code','')
  if not code:
   if 'API' in d['title'] or 'JSON' in d['title']:code='payload = valider(requete)\nscore = modele.predict_proba(payload)\nreponse = appliquer_politique(score)'
   elif 'Pipeline' in d['title'] or 'pipeline' in d['title']:code='model = Pipeline([\n    ("preprocess", preprocessor),\n    ("classifier", LogisticRegression())\n])\nmodel.fit(X_train, y_train)'
   else:code='df = pd.read_csv("novalia_clients.csv")\nX = df.drop(columns=["churn", "customer_id"])\ny = df["churn"]\n\n# Apprendre, puis évaluer sans fuite de données.'
  shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.65,2.0,8.0,4.20,C['navy']);txt(s,code,.91,2.25,7.48,3.7,21,C['white'],minimum=15)
  for i,(t,b) in enumerate(a):card(s,t,b,8.94,2+i*(4.2/len(a)),3.74,4.2/len(a)-.14,ACC[i%6])
 elif k=='sources':
  for i,row in enumerate(it):
   yy=1.97+i*.85;txt(s,row[0],.72,yy,.67,.3,13,C['blue'],True);txt(s,row[1],1.55,yy,10.95,.30,17,C['navy'],True);q=txt(s,row[2],1.55,yy+.33,10.95,.32,11,C['muted']);
   for p in q.text_frame.paragraphs:
    for r in p.runs:
     if str(row[2]).startswith('https://'):r.hyperlink.address=row[2]
 else:grid(s,it)

for day in (1,2):
 prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H);prs.core_properties.title=f'Formation 4-020, Jour {day}';prs.core_properties.author='P2Enjoy Studio';prs.core_properties.subject='IA appliquée aux systèmes d’information de l’entreprise';prs.core_properties.language='fr-FR'
 arr=[s for s in D['slides'] if s['day']==day]
 for d in arr:
  CUR=d['id'];sl=prs.slides.add_slide(prs.slide_layouts[6])
  if d['kind']=='cover':cover(sl,d)
  elif d['kind']=='section':section(sl,d)
  else:frame(sl,d);special(sl,d)
  note=f"{d['id']} | Jour {day} | {d['block']} | Séquence {d['seg']}\nModalité : {d['mode']}. Budget de la séquence complète : {d['segment_minutes']} min (ne pas additionner ce temps à chaque diapositive).\n\nOBJECTIF\n{d['objective']}\n\nDÉVELOPPEMENT ET CONSIGNES FORMATEUR\n{d['n']}\n\nMESSAGE À RETENIR\n{d['take']}"
  for label,key in [('QUESTION DE RELANCE','q'),('RÉPONSE ATTENDUE','a'),('POINT DE VIGILANCE','pitfall')]:
   if d.get(key):note+=f'\n\n{label}\n'+d[key]
  if d.get('sources'):
   note+='\n\nSOURCES\n'+'\n'.join(f"[{r}] {D['sources'][r][0]}\n{D['sources'][r][1]}\n{D['sources'][r][3]}" for r in d['sources'] if r in D['sources'])
  sl.notes_slide.notes_text_frame.text=clean(note)
 prs.save(OUT/f'Formation_4_020_Jour_{day}.pptx')
 print('Created',day,len(arr),flush=True)
Path(os.environ.get('PRODUCTION_TEMP','/mnt/data/_production_4020')).joinpath('text_boxes.json').write_text(json.dumps(META,ensure_ascii=False))
