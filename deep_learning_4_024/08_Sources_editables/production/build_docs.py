from pathlib import Path
import json,re,shutil,textwrap,collections
import numpy as np
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from sklearn.feature_extraction.text import TfidfVectorizer
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'02_Documents';OUT.mkdir(exist_ok=True)
S=json.loads((ROOT/'08_Sources_editables/slides.json').read_text());REF=json.loads((ROOT/'08_Sources_editables/references.json').read_text());C=json.loads((ROOT/'08_Sources_editables/conducteur.json').read_text())
SRC=(ROOT/'08_Sources_editables'/'Dossier_original_fourni.txt').read_text()
# Le dossier source est conservé dans les sources éditables.
BLUE='23468C';GREEN='238C33';RED='F24141';NAVY='122441';GRAY='56677F';PALE='EFF4FA';INK='0D0D0D'
CITEMAP={'turn19search11':'R00','turn3search0':'R23','turn20search1':'R01','turn2search2':'R02','turn4academia19':'R10','turn7search0':'R03','turn4search2':'R11','turn20academia12':'R12','turn7search1':'R04','turn20search3':'R09','turn1view0':'R08','turn4academia18':'R13','turn5academia20':'R14','turn20search2':'R15','turn2search3':'R01','turn2search5':'R07','turn21search1':'R06','turn21search4':'R23','turn23view3':'R16','turn23view4':'R16','turn23view5':'R17','turn26view0':'R19','turn25search3':'R18','turn25search6':'R18','turn22view3':'R20','turn22view5':'R21','turn15search4':'R22','turn21search2':'R02'}
def clean(t):
 def cite(m):
  vals=list(dict.fromkeys(CITEMAP.get(v,'R00') for v in re.findall(r'turn\d+\w+\d+',m.group())))
  return ' ['+', '.join(vals)+']'
 t=re.sub(r'cite.*?',cite,str(t));t=t.replace('—',':').replace('–',' à ').replace('−','-')
 t=t.replace('violation de données','atteinte à la confidentialité des données').replace('violation','non-respect')
 t=t.replace('\\(','').replace('\\)','').replace('\\[','').replace('\\]','')
 formulas={r'z = w_1x_1 + w_2x_2 + \\dots + w_nx_n + b':'z = w₁x₁ + w₂x₂ + … + wₙxₙ + b'}
 t=t.replace('z = w_1x_1 + w_2x_2 + \\dots + w_nx_n + b','z = w₁x₁ + w₂x₂ + … + wₙxₙ + b')
 t=t.replace('\\theta_{t+1}=\\theta_t-\\eta \\nabla_\\theta L','paramètres suivants = paramètres actuels - taux × gradient de la perte')
 t=t.replace('Precision = \\frac{TP}{TP+FP}','Précision = VP / (VP + FP)').replace('Recall = \\frac{TP}{TP+FN}','Rappel = VP / (VP + FN)')
 t=t.replace('\\hat y','y prédit').replace('\\theta','paramètres').replace('\\eta','taux').replace('\\nabla','gradient').replace('\\qquad',' ; ').replace('\\dots','…').replace('\\rightarrow','→')
 for a,b in [('x_{t-k:t}','fenêtre passée'),('y_{t+1}','valeur suivante'),('x_1','x₁'),('x_2','x₂'),('x_n','xₙ'),('w_1','w₁'),('w_2','w₂'),('w_n','wₙ'),('x_i','xᵢ'),('w_i','wᵢ'),('x_t','xₜ'),('h_{t-1}','hₜ₋₁'),('c_t','cₜ'),('^2','²')]:t=t.replace(a,b)
 return t

def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);tcPr.append(sh)

def base(title,subtitle):
 d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.75);sec.right_margin=Cm(1.75);sec.header_distance=Cm(.65);sec.footer_distance=Cm(.75)
 for name in ['Normal','Body Text','List Paragraph']:
  st=d.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(10);st.font.color.rgb=RGBColor.from_string(INK);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.12
 for name,size in [('Title',30),('Subtitle',14),('Heading 1',21),('Heading 2',16),('Heading 3',12)]:
  st=d.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(BLUE);st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
 for name,size,color in [('Métadonnée',8,GRAY),('Légende',8,GRAY),('Code',8,NAVY)]:
  st=d.styles.add_style(name,1);st.font.name='DejaVu Sans Mono' if name=='Code' else 'DejaVu Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.06
 d.styles['Code'].paragraph_format.left_indent=Cm(.25)
 h=sec.header.paragraphs[0];h.text='P2ENJOY   /   FORMATION 4-024   /   DEEP LEARNING PAR LA PRATIQUE';h.style=d.styles['Métadonnée']
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=f.add_run('https://p2enjoy.studio   •   ');r.font.name='DejaVu Sans';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(GRAY)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 d.core_properties.author='Martino Bettucci | P2Enjoy';d.core_properties.title=title;d.core_properties.subject='Formation DGFiP, 21 au 23 septembre 2026, Asnières-sur-Seine';d.core_properties.language='fr-FR'
 d.add_paragraph('FORMATION PROFESSIONNELLE • RÉFÉRENCE 4-024','Subtitle');d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle')
 d.add_paragraph('Intelligence artificielle : Deep Learning par la pratique\n18 heures • 3 jours • Niveau Application (A)\nDGFiP • 21 au 23 septembre 2026 • Asnières-sur-Seine')
 d.add_picture(str(ROOT/'05_Illustrations/001_hero.png'),width=Cm(17.2))
 d.add_paragraph('[MÉTA-DOSSIER | Langue : français | Référence : 4-024 | Édition de production : 14 septembre 2026]','Métadonnée')
 d.add_paragraph('Auteur et formateur : Martino Bettucci • P2Enjoy\nLes exemples métiers du kit sont synthétiques. Les cas institutionnels sont distingués des simulations pédagogiques. Aucune donnée fiscale personnelle n’est utilisée.')
 d.add_page_break();return d

def rich(d,t,style=None):
 p=d.add_paragraph(style=style);parts=re.split(r'(\*\*.*?\*\*|`[^`]+`)',clean(t))
 for part in parts:
  if part.startswith('**') and part.endswith('**'):r=p.add_run(part[2:-2]);r.bold=True
  elif part.startswith('`') and part.endswith('`'):r=p.add_run(part[1:-1]);r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
  else:p.add_run(part)
 return p

def table(d,rows,headers=True):
 t=d.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Light Shading Accent 1'
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,val in enumerate(row):
   cells[j].text=clean(re.sub(r'\*\*|`','',str(val)));cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cells[j].paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
    for r in p.runs:r.font.name='DejaVu Sans';r.font.size=Pt(8.5)
   if i==0 and headers:
    shade(cells[j],BLUE)
    for r in cells[j].paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
  trPr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0 and headers:rep=OxmlElement('w:tblHeader');trPr.append(rep)
 d.add_paragraph('');return t

def h(d,title,level=1,meta=None):
 d.add_heading(clean(title),level)
 if meta:d.add_paragraph(clean(meta),'Métadonnée')

def bibliography(d):
 h(d,'Références et accès aux sources',1,'[MÉTA | Bibliographie | Sources documentaires et documentations officielles]')
 rich(d,'Les identifiants R00 à R23 relient le dossier, les diapositives et les notebooks. Les faits historiques sont rapportés au protocole et à la date de la source ; ils ne constituent pas des performances promises pour les TP. La formulation initiale et ses anciennes balises sont conservées dans le fichier source original.')
 for key,r in REF.items():
  if key=='R24':continue
  rich(d,f"**[{key}] {r['author']}**. {r['title']}. {r['date']}.")
  rich(d,r['url']);rich(d,r['note'])

def conductor(d):
 h(d,'Conducteur opérationnel des 18 heures',1,'[MÉTA | Préparation formateur | Conducteur recalculé | Hors pauses]')
 rich(d,'Ce conducteur remplace les estimations horaires dispersées dans la note de recherche initiale. Il totalise **1 080 minutes**, dont **670 minutes de productions pratiques (11 h 10, soit 62,0 %)**. Les exposés, démonstrations observées, quiz et restitutions ne sont pas comptés comme des travaux pratiques. Les pauses et le déjeuner sont à ajouter aux horaires pédagogiques relatifs.')
 table(d,[['Jour','Pratique','Autres activités','Total'],['1','210 min (3 h 30)','150 min','360 min'],['2','215 min (3 h 35)','145 min','360 min'],['3','245 min (4 h 05)','115 min','360 min']])
 rich(d,'**Légende :** P = production observable par les apprenants ; D = démonstration du formateur ; E = exposé, échange, diagnostic, évaluation ou restitution. Les preuves indiquées doivent être collectées ou contrôlées pendant l’activité.')
 for day in [1,2,3]:
  h(d,f'Jour {day} : six heures pédagogiques',2,f'[MÉTA | J{day} | Temps cumulé : 00:00 à 06:00]')
  rows=[['Temps','Activité','Type / min','Preuve']]
  for x in C:
   if x['day']==day:
    hh=lambda m:f'{m//60:02}:{m%60:02}'
    rows.append([hh(x['start'])+' à '+hh(x['end']),x['label'],x['nature']+' / '+str(x['minutes']),x['proof']])
  table(d,rows)
 rich(d,'**Gestion des équipes :** le créneau de soutenance de 20 minutes correspond à quatre équipes (3 minutes de présentation et 2 minutes de questions chacune). Au-delà de quatre équipes, utiliser des jurys ou restitutions en parallèle plutôt que déborder sur l’épreuve individuelle. Le nombre de groupes doit être fixé avant la session.')

# Choix explicite d’un schéma existant pour remplacer chaque instruction visuelle du dossier.
vec=TfidfVectorizer(ngram_range=(1,2),strip_accents='unicode');mat=vec.fit_transform([s['title']+' '+s['detail']+' '+s['visual'] for s in S]);FIGS=[]
def illustration(d,desc):
 choices=json.loads((ROOT/'08_Sources_editables'/'Choix_figures_dossier.json').read_text());n=len(FIGS)+1;choice=choices['map'][n-1]
 if isinstance(choice,int):
  s=next(v for v in S if v['id']==choice);img=ROOT/'05_Illustrations'/f"{s['id']:03d}_{s['visual']}.png";title=s['title'];sid=choice
 else:
  entry=choices['special'][str(n)];img=ROOT/'05_Illustrations'/(entry['name']+'.png');title=entry['title'];sid=None
 p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(img),width=Cm(16.8))
 d.add_paragraph(f"Figure {n}. {title}. Illustration pédagogique accompagnant le contenu.",'Légende')
 FIGS.append({'figure':n,'instruction_source':desc,'schema':img.name,'diapositive':sid})

def source_body(d):
 lines=SRC.splitlines();i=0;code=False;buf=[]
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('# '):continue
  if line.startswith('```'):
   code=not code;continue
  if code:
   p=d.add_paragraph(line,'Code');p.paragraph_format.keep_with_next=True;continue
  if 'ILLUSTRATION:' in line:
   illustration(d,re.sub(r'^`?ILLUSTRATION:\s*','',line).strip('`'));continue
  if line.startswith('##'):
   level=min(3,len(line)-len(line.lstrip('#')));t=line.lstrip('# ').strip();m=re.search(r'\[MÉTA[^]]*\]',t);meta=m.group() if m else '[MÉTA | Dossier pédagogique | Contenu de référence]';t=re.sub(r'`?\[MÉTA[^]]*\]`?','',t).strip(' `')
   if 'Note méthodologique sans' in t:continue
   if 'Jour-un' in meta:meta='[MÉTA | J1 | 6 h | 3 h 30 de productions pratiques selon le conducteur]'
   if 'Jour-deux' in meta:meta='[MÉTA | J2 | 6 h | 3 h 35 de productions pratiques selon le conducteur]'
   if 'Jour-trois' in meta:meta='[MÉTA | J3 | 6 h | 4 h 05 de productions pratiques selon le conducteur]'
   h(d,t,max(1,level-1),meta);continue
  if line.startswith('|'):
   rows=[]
   while True:
    vals=[x.strip() for x in line.strip('|').split('|')]
    if not all(re.match(r'^:?-+:?$',v.replace(' ','')) for v in vals):rows.append(vals)
    if i>=len(lines) or not lines[i].strip().startswith('|'):break
    line=lines[i].strip();i+=1
   if rows and rows[0][0].startswith(('Temps pédagogique','Temps')):
    rich(d,'**Organisation minutée :** utiliser le conducteur opérationnel situé au début du présent document. Les horaires de la note de recherche ont été remplacés par un décompte vérifié de 6 heures par journée.')
   elif rows:
    n=max(len(r) for r in rows);rows=[r+['']*(n-len(r)) for r in rows];table(d,rows)
   continue
  if any(t in line for t in ['La proposition ci-dessous réserve','La répartition proposée est de','La proposition consacre']):
   if 'Jour 1' in line or '3 h 40' in line:rich(d,'Le Jour 1 conserve les trois blocs contractuels de deux heures. Le conducteur opérationnel réserve 3 h 30 de productions pratiques sur les 6 heures.')
   elif '1 h 45' in line:rich(d,'Le conducteur opérationnel réserve 3 h 35 de productions pratiques au Jour 2, au sein de deux blocs contractuels de trois heures.')
   else:rich(d,'Le conducteur opérationnel réserve 4 h 05 de productions pratiques au Jour 3. Le total vérifié du stage est de 11 h 10 de pratique sur 18 heures (62,0 %).')
   continue
  if line.startswith('`[MÉTA'):
   m=re.match(r'`(\[MÉTA[^`]*\])`(.*)',line)
   if m:d.add_paragraph(clean(m.group(1)),'Métadonnée');line=m.group(2).strip()
   if not line:continue
  if line in ['\\[','\\]','puis']:continue
  if line.startswith('>'):line=line.lstrip('> ')
  if line.startswith('- '):rich(d,'• '+line[2:])
  elif re.match(r'^\d+\. ',line):rich(d,line)
  else:rich(d,line)

QUESTIONS=[
('Un scaler doit être ajusté…',['avant toute séparation','sur l’entraînement uniquement','sur entraînement et test','séparément sur chaque partition'],1,'Ses paramètres doivent être appris sur train puis appliqués sans réajustement aux autres partitions. [R03]'),
('Que réalise loss.backward() ?',['Il met les poids à jour','Il calcule les gradients','Il choisit le seuil','Il sauvegarde le modèle'],1,'La mise à jour est réalisée par optimizer.step(). La rétropropagation calcule les dérivées. [R02]'),
('Avec 99,5 % de négatifs, une exactitude de 99,5 %…',['prouve l’utilité','prouve la calibration','peut être obtenue sans détecter aucun positif','suffit à fixer le seuil'],2,'Un classifieur toujours négatif atteint déjà ce score. Examiner rappel, précision et capacité de traitement. [R04]'),
('Train 99 %, validation 71 % : première hypothèse ?',['Surapprentissage ou décalage à examiner','Conformité garantie','Modèle nécessairement trop petit','Test inutile'],0,'Le contraste suggère un problème de généralisation ; vérifier d’abord le split et les distributions.'),
('Pour CrossEntropyLoss multiclasse, fournir…',['des probabilités et des labels réels','des logits et des indices de classe entiers','un booléen par lot','une sortie sigmoïde unique'],1,'L’API combine la transformation appropriée aux logits avec la perte de classification. [R02]'),
('L’intérêt central d’un CNN sur image est…',['d’apprendre sans données','d’exploiter le voisinage spatial avec filtres partagés','de garantir l’équité','d’éviter toute évaluation'],1,'Localité et partage des paramètres constituent un biais inductif adapté à certaines structures spatiales.'),
('Un LSTM…',['est toujours supérieur à la persistance','maintient des états et des portes différentiables','supprime toutes les fuites temporelles','garantit la prévision financière'],1,'La structure aide à traiter des dépendances séquentielles ; sa pertinence doit être testée. [R09]'),
('Le seuil de revue doit être choisi…',['sur le test après chaque essai','au hasard','sur validation selon coûts et capacité','toujours à 0,5'],2,'Le seuil appartient au système de décision ; il doit être figé avant l’évaluation finale. [R04]'),
('Un transfert prudent commence généralement par…',['tout dégeler avec un taux élevé','geler l’extracteur et entraîner une nouvelle tête','supprimer la validation','réutiliser le test pour entraîner'],1,'L’ajustement de la tête précède un éventuel dégel partiel avec un taux adapté. [R06]'),
('Une supervision humaine effective suppose…',['un simple clic obligatoire','que le score décide seul','l’accès aux éléments utiles et le pouvoir de contester','l’absence de suivi'],2,'Le contrôle doit être réel, informé et permettre de ne pas suivre la sortie du modèle. [R20]'),
]

def evaluations(d,answers=True):
 h(d,'Dispositif d’évaluation et preuves',1,'[MÉTA | Évaluation | Niveau Application | Proposition opérationnelle du kit]')
 rich(d,'Le diagnostic initial n’est pas noté. Les TP sont évalués par des preuves observables. La proposition de validation suivante est pédagogique et ne remplace pas les règles contractuelles de l’organisme : au moins 7/10 au QCM, au moins 60/100 au projet, identification de trois défauts sur cinq et correction de deux défauts à l’épreuve individuelle. Une fuite de test non corrigée impose une remédiation, même si le score global est élevé.')
 h(d,'Diagnostic d’entrée',2,'[MÉTA | J1 | 15 minutes intégrées au conducteur]')
 for q in ['Quelle différence entre validation et test ?','Pourquoi normaliser certaines variables ?','Qu’est-ce qu’un paramètre appris ?','Que signifie une exactitude de 95 % ?','À quoi sert le gradient ?']:rich(d,q+'\nRéponse : __________________________________________________________')
 h(d,'QCM final : une réponse par question',2,'[MÉTA | J3 | 10 minutes | 1 point par réponse correcte]')
 for i,(q,opts,ans,just) in enumerate(QUESTIONS,1):
  rich(d,f'**{i}. {q}**')
  for k,t in enumerate(opts):rich(d,f'{chr(65+k)}. {t}')
  if answers:rich(d,f'**Corrigé : {chr(65+ans)}.** {just}')
 h(d,'Épreuve pratique individuelle',2,'[MÉTA | J3 | 20 minutes | Diagnostic et corrections]')
 rich(d,'À partir du notebook 08, relever cinq défauts : prétraitement avant split ; Dropout actif pendant une évaluation ; sortie/perte ou type de cible incohérents ; exactitude utilisée seule sur une classe rare ; choix répétés après consultation du test. Identifier au moins trois défauts, en corriger deux et expliciter les effets attendus. Le formateur fournit l’extrait défectueux, pas le corrigé.')
 table(d,[['Preuve','Attendu','Vérification'],['Repérage','Au moins 3 défauts sur 5','Annotation du code'],['Correction','Au moins 2 corrections pertinentes','Cellules exécutées / assertions'],['Justification','Conséquence sur apprentissage ou évaluation','3 phrases explicites']])
 h(d,'Grille du projet final',2,'[MÉTA | Équipe | Barème /100 | SAME Application]')
 rows=[['Critère','Points','Preuve attendue']]+[[a,str(b),c] for a,b,c in [('Problème métier',15,'Cible, utilisateur, moment de prédiction'),('Données et séparation',15,'Partitions disjointes, preprocessing sur train'),('Baseline',10,'Référence simple pertinente'),('Architecture',10,'Choix cohérent avec les entrées'),('Protocole expérimental',15,'Hypothèse et essai documentés'),('Métriques',15,'Mesure adaptée et seuil justifié'),('Analyse d’erreurs',10,'Trois erreurs ou limites examinées'),('Risques et recommandation',10,'Usage limité, contrôle humain et prochaine étape')]]
 table(d,rows)
 rich(d,'**Le score brut du réseau ne rapporte pas de points à lui seul.** Une équipe concluant de manière fondée à l’insuffisance de son modèle peut réussir l’évaluation.')
 h(d,'Remédiations',2,'[MÉTA | Formateur | Ajustement au niveau observé]')
 table(d,[['Difficulté','Intervention','Nouvelle preuve'],['Python ou formes','Cellules guidées et schéma du tenseur','Assertion corrigée'],['Fuite de données','Reconstituer le moment réel de prédiction','Nouveau split et scaler'],['Métriques','Revenir à 100 cas avec classe rare','Confusion et charge expliquées'],['Diagnostic','Limiter à une hypothèse et un test','Journal renseigné'],['Interprétation excessive','Distinguer observation et recommandation','Conclusion métier bornée']])
 h(d,'Traçabilité objectifs, activités et évaluations',2,'[MÉTA | Qualité pédagogique | Preuves conservées]')
 table(d,[['Compétence','Activités','Preuves / évaluation'],['Comprendre les réseaux','J1, TP 01','Gradient, formes, QCM 2 et 6'],['Utiliser les bibliothèques','J1, TP 00 et 02','Imports, entraînement, sauvegarde'],['Concevoir un CNN et un LSTM','J2, TP 04 et 06','Dimensions, modèles, comparaison'],['Diagnostiquer et optimiser','J2, TP 05, 07, 08','Journal, corrections, épreuve individuelle'],['Appliquer à un métier','J3, TP 09 et 10','Fiche de cadrage, projet /100'],['Documenter les limites','Cas et projet','Fiches modèle/système, soutenance']])
 h(d,'Suivi de réalisation et retour de formation',2,'[MÉTA | Administration de formation | Modèles à renseigner]')
 rich(d,'Recueillir les présences par demi-journée, les productions significatives, les résultats d’évaluation et le questionnaire de satisfaction. Le certificat de réalisation atteste la réalisation de l’action ; ne pas le présenter comme une certification professionnelle de maîtrise.')
 table(d,[['Indicateur','Calcul proposé','À conserver'],['Satisfaction','Moyenne des réponses valides sur échelle annoncée','Effectif répondant et distribution'],['Nombre d’apprenants','Nombre de participants effectivement accueillis','Liste et présences'],['Abandons','Abandons / inscrits ayant commencé','Motif, date, suivi'],['Taux de retour','Questionnaires reçus / questionnaires attendus','Dénominateur explicite'],['Interruptions','Prestations interrompues / prestations commencées','Cause et durée']])


def main():
 d=base('Dossier pédagogique\net guide de préparation','Contenus développés, cas documentés, conducteur, illustrations, évaluation')
 h(d,'Repères de lecture et périmètre',1,'[MÉTA | Note éditoriale | Mise en production du dossier fourni]')
 rich(d,'Ce document met en page le dossier fourni et remplace ses consignes visuelles par des schémas effectivement produits. Le fichier original est conservé sans modification dans les sources éditables. Les 126 diapositives comportent des notes orateur, des questions d’animation, les réponses attendues et des références. Les notebooks apprenants et corrigés sont distribués séparément.')
 rich(d,'**Adaptations explicites de production :** le minutage a été recalculé pour respecter les 18 heures et le minimum de pratique ; les exercices ont une voie CPU avec des données locales ; le laboratoire vision emploie des chiffres 8 × 8 plutôt qu’un téléchargement d’images 28 × 28 ; le transfert est simulé à petite échelle sur des sous-ensembles distincts. Keras est utilisé avec le moteur PyTorch dans la recette locale. La branche native TensorFlow reste conditionnelle et son absence est signalée dans le rapport de recette. Les formules sont présentées en typographie classique.')
 h(d,'Sommaire',2,'[MÉTA | Navigation | Les titres Word sont structurés pour le volet de navigation]')
 for t in ['Conducteur opérationnel des 18 heures','Rappels indispensables','Histoire et architectures','Environnement et bibliothèques','Jour 1 : premier modèle','Jour 2 : conception et diagnostic','Jour 3 : cas et projet','Gouvernance et conformité','Évaluation et kit formateur','Évaluations opérationnelles et corrigés','Références documentaires']:rich(d,t)
 conductor(d);d.add_page_break();source_body(d);evaluations(d,True);bibliography(d)
 path=OUT/'Dossier_pedagogique_4_024.docx';d.save(path);print(path.name,'paragraphes',len(d.paragraphs),'figures',len(FIGS),flush=True)
 (ROOT/'08_Sources_editables'/'Correspondance_illustrations.json').write_text(json.dumps(FIGS,ensure_ascii=False,indent=2))
 # Guide orateur : toutes les explications, questions et corrections, sans dupliquer le dossier entier.
 g=base('Guide formateur\nNotes des 126 diapositives','Animation, questions, réponses attendues et points de vigilance')
 rich(g,'Les numéros correspondent exactement au diaporama global. Les trois fichiers journaliers reprennent respectivement les diapositives 001 à 042, 043 à 084 et 085 à 126. Les notes sont également intégrées à chaque diapositive PowerPoint.')
 conductor(g)
 for day in [1,2,3]:
  g.add_page_break();h(g,f'Jour {day} : notes d’animation',1,f'[MÉTA | J{day} | Guide formateur | Numérotation globale]')
  for s in [v for v in S if v['day']==day]:
   h(g,f"{s['id']:03d}. {s['title']}",2,s['meta']);rich(g,'**Message central :** '+s['lead']);rich(g,s['detail']);rich(g,'**Question à poser :** '+s['question']);rich(g,'**Réponse attendue :** '+s['answer']);rich(g,'**Preuve à rechercher :** une reformulation, un calcul, une assertion ou une observation issue du notebook. Distinguer le fait observé de son interprétation.');g.add_paragraph('Sources : '+', '.join(s['refs']),'Métadonnée')
 bibliography(g);g.save(OUT/'Guide_formateur_126_diapos.docx');print('Guide formateur créé',flush=True)
 # Cahier de travaux : les consignes originales de tous les notebooks avec des espaces de travail.
 w=base('Cahier participant\nTravaux pratiques et projet','11 ateliers, consignes, preuves attendues et fiches à compléter')
 rich(w,'Ce cahier accompagne les notebooks du dossier 03_Travaux_pratiques/apprenants. Il ne contient pas les solutions. Lire chaque consigne avant d’exécuter le code, conserver les essais et remettre les productions demandées. Les données métiers sont synthétiques ; aucune conclusion concernant une personne réelle n’est permise.')
 table(w,[['Jour','Ateliers principaux','Preuves'],['1','00, 01, 02, 03','Environnement, split, entraînement, métriques'],['2','04, 05, 06, 07, 08','CNN, séquences, diagnostic, transfert'],['3','09, 10','Capacité, dérive, projet justifié']])
 for nbp in sorted((ROOT/'03_Travaux_pratiques/apprenants').glob('*.ipynb')):
  w.add_page_break();nb=json.loads(nbp.read_text());h(w,nbp.name[:2]+' : '+nb['cells'][0]['source'][0].replace('# ','').strip(),1,f'[MÉTA | TP {nbp.name[:2]} | Participant | Notebook : {nbp.name}]')
  for c in nb['cells']:
   if c['cell_type']!='markdown':continue
   text=''.join(c['source']) if isinstance(c['source'],list) else c['source']
   for line in text.splitlines():
    if line.startswith('# '):continue
    if line.startswith('## '):h(w,line[3:],2)
    elif line.startswith('`[MÉTA'):continue
    elif line.strip():rich(w,line)
  h(w,'Journal d’expérience à remplir',2,'[MÉTA | Preuve individuelle ou d’équipe]')
  table(w,[['Hypothèse','Modification','Métrique attendue','Observation','Décision'],['________________','________________','________________','________________','________________'],['________________','________________','________________','________________','________________']])
  rich(w,'**Interprétation :** ce que le résultat permet d’affirmer, ce qu’il ne permet pas d’affirmer et la prochaine vérification utile.\n\n________________________________________________________________________\n\n________________________________________________________________________')
 h(w,'Fiche de données et de modèle',1,'[MÉTA | Projet final | À rendre avec le notebook]')
 table(w,[['Champ','Réponse'],*[[x,'___________________________________________________'] for x in ['Finalité et moment de prédiction','Utilisateur et action autorisée','Provenance et période des données','Cible et méthode d’annotation','Variables retenues / exclues','Partitions et prétraitements','Référence simple et architecture','Métrique principale et seuil','Erreurs, sous-groupes et limites','Conditions de supervision','Surveillance et conditions d’arrêt']]])
 for par in w.paragraphs:
  if par.text=='Journal d’expérience à remplir':par.paragraph_format.page_break_before=True
  if not par.text.strip() and not par._p.xpath('.//w:br'):par.paragraph_format.keep_with_next=True
  if par.style.name=='Heading 1' and re.match(r'^\d\d : \d\d',par.text):par.text=re.sub(r'^(\d\d) : \d\d\s*[•:]\s*',r'TP \1 : ',par.text)
  if par.text.startswith('Interprétation :'):par.paragraph_format.keep_together=True
 for tt in w.tables:
  if tt.cell(0,0).text=='Hypothèse':
   for rr in tt.rows:
    for cc in rr.cells:
     for pp in cc.paragraphs:pp.paragraph_format.keep_with_next=True
 w.save(OUT/'Cahier_participant_TP_et_projet.docx');print('Cahier participant créé',flush=True)
 e=base('Évaluations\net corrigés formateur','Diagnostic, QCM, épreuve pratique et grille du projet')
 evaluations(e,True);e.save(ROOT/'04_Evaluations'/'Evaluations_et_corriges.docx')
 # Version indépendante distribuable sans réponses.
 e2=base('Évaluations\nVersion participant','Diagnostic, QCM, épreuve pratique et critères de réussite')
 evaluations(e2,False);e2.save(ROOT/'04_Evaluations'/'Evaluations_participant.docx')
 (ROOT/'04_Evaluations'/'QCM_et_corriges.json').write_text(json.dumps([{'numero':i+1,'question':q,'choix':o,'reponse':chr(65+a),'justification':j} for i,(q,o,a,j) in enumerate(QUESTIONS)],ensure_ascii=False,indent=2))
 print('Évaluations créées',flush=True)

def apply(d):
 d.styles['Métadonnée'].paragraph_format.keep_with_next=True
 ps=d.paragraphs
 for i,p in enumerate(ps):
  if p.style.name!='Normal' or not re.match(r'^\d+\. ',p.text):continue
  opts=[]
  for j in range(i+1,min(i+6,len(ps))):
   if re.match(r'^[A-D]\. ',ps[j].text):opts.append(j)
   else:break
  if len(opts)==4:
   p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
   for j in opts:
    ps[j].paragraph_format.keep_together=True;ps[j].paragraph_format.keep_with_next=(j!=opts[-1] or (j+1<len(ps) and ps[j+1].text.startswith('Corrigé')))
   if opts[-1]+1<len(ps) and ps[opts[-1]+1].text.startswith('Corrigé'):ps[opts[-1]+1].paragraph_format.keep_together=True
if __name__=='__main__':
 main()
 for pp in list(OUT.glob('*.docx'))+list((ROOT/'04_Evaluations').glob('*.docx')):
  dd=Document(pp);apply(dd);dd.save(pp)

