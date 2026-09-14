from pathlib import Path
from PIL import ImageFont
import json,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
ROOT=Path(__file__).resolve().parents[2];D=ROOT/'01_Presentations';D.mkdir(exist_ok=True)
S=json.loads((ROOT/'08_Sources_editables/slides.json').read_text()); R=json.loads((ROOT/'08_Sources_editables/references.json').read_text())
B='23468C';G='238C33';Y='D9CF4A';RED='F24141';N='122441';BLACK='0D0D0D';GRAY='56677F';P='EFF4FA';WHITE='FFFFFF';FONT='DejaVu Sans'
def clean(t):return str(t).replace('—',':').replace('–','à').replace('violation','atteinte à la protection').replace('inviolable','réputé sûr')
def rect(sl,x,y,w,h,fill,line=None,r=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background() if line is None else None
 if line:sh.line.color.rgb=RGBColor.from_string(line)
 if r:
  try:sh.adjustments[0]=.09
  except Exception:pass
 return sh

def text(sl,x,y,w,h,t,size=20,color=BLACK,bold=False,align=None):
 tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,l in enumerate(clean(t).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.08
  if align is not None:p.alignment=align
 return tb

def notes(s):
 n=clean(s['notes']);n+='\n\nLÉGENDE DE L’ILLUSTRATION\nSchéma de formation. Les valeurs montrées dans les graphiques conceptuels sont illustratives. Les exemples d’images sont issus du jeu local de chiffres manuscrits. Les cas DGFiP ne reproduisent pas des systèmes internes.\n\nRÉFÉRENCES CONSULTABLES\n'
 for rid in s['refs']:
  r=R.get(rid)
  if r:n+=f"[{rid}] {r['author']}. {r['title']}. {r['date']}. {r['url']}\n"
 return n

def build(slides,name):
 p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5);p.core_properties.title='Intelligence artificielle : Deep Learning par la pratique';p.core_properties.author='Martino Bettucci | P2Enjoy';p.core_properties.subject='Formation 4-024, DGFiP, 21 au 23 septembre 2026';p.core_properties.language='fr-FR';p.core_properties.keywords='Deep Learning, 4-024, DGFiP, formation'
 for s in slides:
  sl=p.slides.add_slide(p.slide_layouts[6]);accent={1:B,2:G,3:RED}[s['day']];v=ROOT/'05_Illustrations'/f"{s['id']:03d}_{s['visual']}.png"
  cover=s['kind']=='cover';section=s['kind']=='section'
  if cover:
   sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(N)
   rect(sl,0,0,.19,7.5,accent);text(sl,.55,.35,9,.35,'P2ENJOY  /  FORMATION PROFESSIONNELLE  /  RÉF. 4-024',11,WHITE,True)
   text(sl,.6,1.12,6.5,1.7,s['title'],39,WHITE,True);text(sl,.62,3.04,5.8,.86,s['lead'],22,WHITE)
   for i,item in enumerate(s['items']):text(sl,.64,4.24+i*.47,5.8,.42,item,17,WHITE)
   rect(sl,6.65,1.72,6.18,3.36,WHITE,r=True);sl.shapes.add_picture(str(v),Inches(6.81),Inches(1.91),width=Inches(5.87),height=Inches(2.94))
   crop=ROOT/'05_Illustrations/detail_reseau_imagegen.png'
   if crop.exists():sl.shapes.add_picture(str(crop),Inches(10.58),Inches(.41),width=Inches(2.23),height=Inches(.89))
   rect(sl,.63,6.28,12.09,.58,'213B5D',r=True);text(sl,.85,6.4,11.65,.34,s['question'],15,WHITE)
   text(sl,.64,7.1,8,.22,'https://p2enjoy.studio  •  DGFiP  •  Asnières-sur-Seine',9,WHITE)
   text(sl,11.6,7.03,1.09,.32,f"JOUR {s['day']}",12,WHITE,True,PP_ALIGN.RIGHT)
  else:
   sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
   rect(sl,0,0,13.333,.11,accent)
   rect(sl,.54,.35,.83,.32,accent,r=True);text(sl,.58,.405,.75,.22,f"JOUR {s['day']}",9,WHITE,True,PP_ALIGN.CENTER)
   typ={'activity':'MISE EN PRATIQUE','concept':'REPÈRE ET MÉTHODE','section':'NOUVELLE SÉQUENCE'}.get(s['kind'],'REPÈRE')
   text(sl,1.53,.40,9,.3,f"{s['block']}  /  {typ}",11,GRAY,True);text(sl,10.25,.39,2.52,.3,'P2ENJOY.STUDIO',11,accent,True,PP_ALIGN.RIGHT)
   fontsize=31
   while ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',fontsize*4).getlength(s['title'])/4 > 12.0*72 and fontsize>22:fontsize-=1
   text(sl,.56,.94,12.18,.97,s['title'],fontsize,N,True)
   text(sl,.6,1.88,12.1,.54,s['lead'],18,GRAY)
   for i,item in enumerate(s['items'][:3]):
    yy=2.62+i*1.03;rect(sl,.6,yy,4.50,.91,P,r=True);rect(sl,.6,yy,.055,.91,accent);text(sl,.80,yy+.16,4.05,.65,item,18.5,BLACK)
   rect(sl,5.33,2.47,7.4,3.92,WHITE,'D8E1ED',True);sl.shapes.add_picture(str(v),Inches(5.47),Inches(2.60),width=Inches(7.12),height=Inches(3.56))
   qlabel='À FAIRE' if s['kind']=='activity' else 'À DISCUTER';rect(sl,.6,6.47,12.14,.56,N,r=True);text(sl,.80,6.62,1.48,.23,qlabel,10.5,Y,True);text(sl,2.33,6.59,10.18,.35,s['question'],13.5,WHITE)
   text(sl,.61,7.21,10.50,.18,f"4-024  •  {', '.join(s['refs'])}  •  Sources complètes et explications dans les notes",8.5,GRAY)
   text(sl,11.32,7.16,1.43,.27,f"{s['id']:03d} / 126",10.5,accent,True,PP_ALIGN.RIGHT)
  sl.notes_slide.notes_text_frame.text=notes(s)
 p.save(D/name)
 print(name,len(slides),flush=True)

if __name__=='__main__':
 names={1:'Jour_1_Fondamentaux_et_premiers_modeles.pptx',2:'Jour_2_CNN_LSTM_et_diagnostic.pptx',3:'Jour_3_Cas_metiers_et_projet.pptx'}
 for d in [1,2,3]:build([s for s in S if s['day']==d],names[d])
 build(S,'Formation_4_024_126_diapos.pptx')
